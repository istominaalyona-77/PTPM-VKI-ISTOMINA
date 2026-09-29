# -*- coding: utf-8 -*-
"""
Генератор отчёта по лабораторной работе №2.

Запуск (из корня репозитория Lab2_Istomina):
    python make_report.py

Требуется один раз установить библиотеку для работы с .docx:
    python -m pip install python-docx

Результат: файл «Отчет_ЛР2_Истомина.docx» в корне репозитория.
Все числовые значения в отчёте соответствуют фактическому выводу
python -m unittest discover -v: Ran 101 tests, FAILED (failures=4, errors=2).
"""

import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(14)

title = doc.add_heading("ОТЧЁТ по лабораторной работе №2", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph("Тема: Юнит-тестирование проектов и автоматизация проверок")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

table = doc.add_table(rows=4, cols=2)
table.style = "Table Grid"
data = [
    ("Студент", "Истомина А. И., группа 2407са2"),
    ("Курс", "3"),
    ("Проверил", "к.т.н., доцент КИ ВКИ НГУ Ярославцева Т. В."),
    ("Новосибирск", "2026"),
]
for row, (key, value) in zip(table.rows, data):
    row.cells[0].text = key
    row.cells[1].text = value

# ---------------------------------------------------------------- 1
doc.add_heading("1. Общие сведения о работе", level=1)
doc.add_paragraph(
    "В ходе лабораторной работы выполнено юнит-тестирование двух модулей: "
    "собственного кода, реализованного в лабораторной работе №1 (проверка данных "
    "пользователя при регистрации), и учебного модуля расчёта стоимости доставки "
    "delivery.py, предоставленного преподавателем."
)
doc.add_paragraph(
    "Тесты разработаны с использованием стандартного модуля unittest. "
    "Запуск выполняется командой «python -m unittest discover -v» из корня репозитория."
)

doc.add_paragraph("Структура проекта:")
for line in [
    "Lab2_Istomina/",
    "├── src/                       исходный код",
    "│   ├── registration.py         модуль из ЛР1 (исправленный)",
    "│   ├── delivery.py             учебный модуль (в исходном виде, с дефектами)",
    "│   └── delivery_fixed.py       исправленная версия для проверки гипотез",
    "├── tests/                     автотесты",
    "│   ├── test_registration.py    50 тестов",
    "│   ├── test_delivery.py        37 тестов",
    "│   └── test_delivery_fixed.py  14 тестов",
    "├── .gitignore",
    "├── make_report.py              генератор настоящего отчёта",
    "└── Отчет_ЛР2_Истомина.docx",
]:
    doc.add_paragraph(line, style="List Paragraph")

# ---------------------------------------------------------------- 2. Пункт А
doc.add_heading("2. Пункт А. Общее количество тестов и статистика прохождения", level=1)
doc.add_paragraph("Всего написано 101 юнит-тест. Сводная статистика:", style="List Bullet")

stat = doc.add_table(rows=5, cols=4)
stat.style = "Table Grid"
stat.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["Модуль", "Тестов", "Успешно", "Не прошло"]):
    stat.rows[0].cells[i].text = h

rows = [
    ("registration.py (свой код, ЛР1)", "50", "50", "0"),
    ("delivery.py (учебный модуль)", "37", "31", "6"),
    ("delivery_fixed.py (исправленная версия)", "14", "14", "0"),
    ("ИТОГО", "101", "95", "6"),
]
for row, values in zip(stat.rows[1:], rows):
    for cell, value in zip(row.cells, values):
        cell.text = value

for cell in stat.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

doc.add_paragraph(
    "Результат выполнения команды «python -m unittest discover -v»: "
    "Ran 101 tests — FAILED (failures=4, errors=2). Все шесть непрошедших проверок "
    "приходятся на учебный модуль delivery.py. После исправления выявленных дефектов "
    "все 101 тест проходят успешно."
)

# ---------------------------------------------------------------- 3. Пункт Б
doc.add_heading("3. Пункт Б. Количество упавших тестов для delivery.py", level=1)
p = doc.add_paragraph()
p.add_run("Количество упавших тестов: 6 (из 37 написанных для модуля delivery.py).").bold = True

doc.add_paragraph(
    "Из них 4 завершились как FAIL (нарушено ожидаемое утверждение) и 2 — как ERROR "
    "(вместо результата выбрасывается необработанное исключение TypeError). "
    "Эти шесть проверок вскрывают три различных дефекта исходного кода."
)

# ---------------------------------------------------------------- 4. Пункт В
doc.add_heading("4. Пункт В. Локализация аномалий", level=1)


def bold_table_head(table_):
    for cell in table_.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def small_font(table_):
    for row in table_.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


doc.add_heading("Дефект 1. Некорректная проверка типов входных параметров", level=2)
doc.add_paragraph("Строка кода: 12")
c1 = doc.add_paragraph()
c1.add_run("if weight < 0.1 or weight > 50.0 or distance < 1 or distance > 5000:").italic = True

doc.add_paragraph("Описание дефекта:", style="List Bullet")
doc.add_paragraph(
    "Строка сразу выполняет арифметическое сравнение, не проверяя типы аргументов. "
    "При передаче нечислового веса (например, строки «тяжёлый») или нечислового расстояния "
    "интерпретатор выбрасывает необработанное исключение TypeError. При этом "
    "документация функции прямо требует возвращать (-1, «0000-00-00») при неверных "
    "параметрах, то есть функция нарушает собственный контракт. Дефект приводит к аварийному "
    "завершению программы вместо штатной обработки ошибочного ввода."
)
doc.add_paragraph("Падающие тесты:", style="List Bullet")
doc.add_paragraph("test_non_numeric_weight_returns_error_instead_of_exception", style="List Bullet 2")
doc.add_paragraph("test_non_numeric_distance_returns_error_instead_of_exception", style="List Bullet 2")

doc.add_paragraph("Предлагаемое исправление:", style="List Bullet")
f1 = doc.add_paragraph()
f1.add_run(
    "Добавить проверку типов перед сравнением:\n"
    "    if not isinstance(weight, (int, float)) or isinstance(weight, bool):\n"
    "        return -1, \"0000-00-00\"\n"
    "    if not isinstance(distance, int) or isinstance(distance, bool):\n"
    "        return -1, \"0000-00-00\""
).italic = True

doc.add_heading("Дефект 2. Экспресс-доставка стоит дешевле обычной", level=2)
doc.add_paragraph("Строка кода: 36")
c2 = doc.add_paragraph()
c2.add_run("total_cost *= 0.5").italic = True

doc.add_paragraph("Описание дефекта:", style="List Bullet")
doc.add_paragraph(
    "При заказе экспресс-доставки итоговая стоимость уменьшается вдвое. Это "
    "противоречит самому смыслу услуги: экспресс обеспечивает более быструю "
    "доставку за счёт привлечения дополнительных ресурсов (авиаперелёт, курьер), "
    "поэтому он должен быть дороже обычного тарифа, а не дешевле. Ошибка в знаке "
    "коэффициента инвертирует экономику доставки: при расстоянии 100 км обычная "
    "доставка стоит 700 рублей, а экспрессная — 350 рублей. Кроме того, скидка 50 % "
    "уничтожает надбавку за особые условия перевозки: опасный груз на 100 км стоит "
    "1700 рублей обычным тарифом и 850 рублей экспрессом. Дефект затрагивает любой "
    "заказ с флагом is_express независимо от веса, расстояния и типа посылки."
)
doc.add_paragraph("Падающие тесты:", style="List Bullet")
doc.add_paragraph("test_express_delivery_should_cost_more_than_standard", style="List Bullet 2")
doc.add_paragraph("test_express_hazardous_should_cost_more_than_standard", style="List Bullet 2")

doc.add_paragraph("Предлагаемое исправление:", style="List Bullet")
f2 = doc.add_paragraph()
f2.add_run(
    "Заменить коэффициент скидки на коэффициент надбавки:\n"
    "    if is_express:\n"
    "        total_cost *= 1.5  # надбавка +50 % за ускоренную доставку"
).italic = True

doc.add_heading("Дефект 3. Экспресс-доставка в день отправки", level=2)
doc.add_paragraph("Строка кода: 44")
c3 = doc.add_paragraph()
c3.add_run("days_needed = days_needed // 2").italic = True

doc.add_paragraph("Описание дефекта:", style="List Bullet")
doc.add_paragraph(
    "Ограничение «минимальный срок доставки — один день», установленное на строке 41 "
    "выражением max(1, distance // 500), применяется ДО деления на два при экспресе. "
    "Целочисленное деление 1 // 2 даёт 0, поэтому посылка, отправленная сегодня, "
    "приходит в тот же день: заказ от 3 сентября доставляется 3 сентября. "
    "Физически это невозможно, а кроме того, нарушается заявленный минимальный срок."
)
doc.add_paragraph("Границы проявления дефекта (проверено перебором всех расстояний):", style="List Bullet")
doc.add_paragraph(
    "Дефект проявляется для всех расстояний от 1 до 999 км включительно — это 999 значений "
    "из 5000 допустимых, то есть около 20 % заказов экспресс-доставки. Граница прохождения "
    "ровно 1000 км вызвана тем, что при ней 1000 // 500 = 2, а 2 // 2 = 1, и минимальный "
    "день сохраняется. Для расстояний 1000 км и более счёт дней всегда положителен, поэтому "
    "дефект себя не проявляет.",
    style="List Bullet 2",
)

boundary = doc.add_table(rows=6, cols=3)
boundary.style = "Table Grid"
for i, h in enumerate(["Расстояние", "Срок экспресса (сейчас)", "Ожидается"]):
    boundary.rows[0].cells[i].text = h
brows = [
    ("1 км", "0 дней — дефект", "не менее 1 дня"),
    ("499 км", "0 дней — дефект", "не менее 1 дня"),
    ("999 км", "0 дней — дефект", "не менее 1 дня"),
    ("1000 км", "1 день", "не менее 1 дня"),
    ("5000 км", "5 дней", "не менее 1 дня"),
]
for row, values in zip(boundary.rows[1:], brows):
    for cell, value in zip(row.cells, values):
        cell.text = value
bold_table_head(boundary)

doc.add_paragraph("Падающие тесты:", style="List Bullet")
doc.add_paragraph("test_express_delivery_should_not_be_same_day_as_shipping", style="List Bullet 2")
doc.add_paragraph("test_express_delivery_time_is_never_zero_days", style="List Bullet 2")

doc.add_paragraph("Предлагаемое исправление:", style="List Bullet")
f3 = doc.add_paragraph()
f3.add_run(
    "Перенести ограничение минимума внутрь условия экспресса:\n"
    "    if is_express:\n"
    "        days_needed = max(1, days_needed // 2)"
).italic = True

# ---------------------------------------------------------------- 5
doc.add_heading("5. Сводная таблица локализации аномалий", level=1)
summary = doc.add_table(rows=7, cols=4)
summary.style = "Table Grid"
for i, h in enumerate(["Упавший тест", "Строка", "Суть дефекта", "Исправление"]):
    summary.rows[0].cells[i].text = h
defects = [
    ("test_non_numeric_weight_returns_error_instead_of_exception", "12",
     "TypeError вместо кода возврата ошибки", "isinstance-проверка типа weight"),
    ("test_non_numeric_distance_returns_error_instead_of_exception", "12",
     "TypeError вместо кода возврата ошибки", "isinstance-проверка типа distance"),
    ("test_express_delivery_should_cost_more_than_standard", "36",
     "Экспресс дешевле обычного тарифа", "total_cost *= 1.5"),
    ("test_express_hazardous_should_cost_more_than_standard", "36",
     "Надбавка за опасный груз поглощается скидкой", "total_cost *= 1.5"),
    ("test_express_delivery_should_not_be_same_day_as_shipping", "44",
     "Доставка в день отправки (срок 0 дней)", "max(1, days_needed // 2)"),
    ("test_express_delivery_time_is_never_zero_days", "44",
     "Срок доставки обнуляется целочисленным делением", "max(1, days_needed // 2)"),
]
for row, values in zip(summary.rows[1:], defects):
    for cell, value in zip(row.cells, values):
        cell.text = value
bold_table_head(summary)
small_font(summary)

# ---------------------------------------------------------------- 6
doc.add_heading("6. Проверка гипотез об исправлениях", level=1)
doc.add_paragraph(
    "Чтобы убедиться в правильности диагноза, исправленный модуль delivery_fixed.py "
    "был реализован в отдельном файле, не затрагивая исходный delivery.py. Для него "
    "написан отдельный набор из 14 тестов, проверяющий как исправления, так и "
    "отсутствие регрессий в ранее корректных сценариях."
)
doc.add_paragraph("Результат: 14 из 14 тестов проходят успешно.", style="List Bullet")
doc.add_paragraph(
    "Дополнительно все 37 тестов из test_delivery.py были запущены против исправленного "
    "модуля (через временную замену импорта) — также 37 из 37 успешно. Это подтверждает, "
    "что диагноз локализован верно: исправления устраняют ровно те аномалии, которые "
    "зафиксированы в отчёте, и не затрагивают корректные сценарии.",
    style="List Bullet",
)

# ---------------------------------------------------------------- 7
doc.add_heading("7. Дефекты, найденные и исправленные в собственном коде (ЛР1)", level=1)
doc.add_paragraph(
    "Тестирование собственного кода из лабораторной работы №1 выявило четыре дефекта, "
    "которые были исправлены. После исправлений все 50 тестов проходят успешно."
)
own = [
    ("1", "registration.py, строка 8",
     "Папка logs создавалась только вручную. При её отсутствии FileHandler "
     "выбрасывал FileNotFoundError, и программа падала при запуске.",
     "Добавлено os.makedirs(\"logs\", exist_ok=True)"),
    ("2", "registration.py, функция mask",
     "При длине пароля 2 символа функция возвращала пароль целиком "
     "(«ab» → «ab»), полностью раскрывая секрет в логах. Это прямое нарушение "
     "требования маскирования конфиденциальных данных.",
     "При длине меньше 3 символов возвращаются только звёздочки"),
    ("3", "registration.py, EMAIL_PATTERN",
     "Регулярное выражение принимало невалидные адреса: «ivan@.ru», "
     "«ivan..petrov@mail.ru», «ivan@mail.ru.».",
     "Добавлена функция is_valid_email с запретом двойных точек"),
    ("4", "registration.py, функция check_password",
     "Пробел и другие символы, не являющиеся буквами, цифрами или спецсимволами, "
     "молча проходили валидацию, так как isalpha() для них возвращает False.",
     "Добавлена явная проверка состава символов пароля"),
]
own_table = doc.add_table(rows=len(own) + 1, cols=4)
own_table.style = "Table Grid"
for i, h in enumerate(["№", "Расположение", "Описание дефекта", "Исправление"]):
    own_table.rows[0].cells[i].text = h
for row, values in zip(own_table.rows[1:], own):
    for cell, value in zip(row.cells, values):
        cell.text = value
bold_table_head(own_table)
small_font(own_table)

# ---------------------------------------------------------------- 8
doc.add_heading("8. Дополнительные замечания", level=1)
doc.add_paragraph(
    "В модуле delivery.py на строке 11 сохранился комментарий «TODO: Проверить "
    "физические ограничения на вес и дистанцию», хотя проверка уже реализована на "
    "строке 12. Это не влияет на работу функции, но устаревший комментарий вводит "
    "в заблуждение и должен быть удалён.",
    style="List Bullet",
)
doc.add_paragraph(
    "Стиль задания границ веса непоследователен: нижняя граница задана строгим "
    "оператором «> 5.0», а верхняя — нестрогим «>= 20.0». Поведение при этом "
    "корректно (вес ровно 5 кг не требует надбавки), поэтому дефектом это не "
    "считается, но единообразие сделает код понятнее.",
    style="List Bullet",
)
doc.add_paragraph(
    "Чёрный список логинов в registration.py применяется только к логинам-строкам: "
    "проверка выполняется после ветвления форматов телефона и email, поэтому адрес "
    "вида admin@mail.ru проходит валидацию. Формально требование задания «логин не "
    "должен совпадать с логинами из чёрного списка» сформулировано без оговорок на "
    "формат, поэтому данное поведение следует считать неоднозначным и уточнить "
    "у мастера практики.",
    style="List Bullet",
)

# ---------------------------------------------------------------- 9
doc.add_heading("9. Заключение", level=1)
doc.add_paragraph(
    "В ходе лабораторной работы написано 101 юнит-тест для трёх модулей. "
    "Все 50 тестов собственного кода из ЛР1 проходят успешно: в ходе работы найдено "
    "и исправлено 4 дефекта, включая утечку короткого пароля в логах и падение "
    "программы при отсутствии папки logs."
)
doc.add_paragraph(
    "Для учебного модуля delivery.py написано 37 тестов, из которых 6 не проходят. "
    "Все они локализованы в трёх строках исходного кода (12, 36, 44) и относятся к трём "
    "независимым дефектам: отсутствию проверки типов, инверсии коэффициента "
    "экспресс-доставки и обнулению срока доставки. Исправления реализованы в "
    "delivery_fixed.py и подтверждены 14 проходящими тестами, а также повторным "
    "прогоном всех 37 тестов против исправленной версии."
)
doc.add_paragraph(
    "Освоенные в ходе работы навыки: разработка юнит-тестов на основе unittest, "
    "применение subTest и параметризации, анализ упавших тестов и локализация "
    "дефектов по строке исходника, проверка гипотез исправлением через "
    "тестирование, а также организация проекта с изоляцией кода в папке src "
    "и тестов в папке tests."
)


path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    'Отчет_ЛР2_Истомина.docx',
)
doc.save(path)
print('Отчёт сохранён:', path)
